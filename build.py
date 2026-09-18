"""Excel -> products.js.  Run after editing the sheet:  python build.py [path.xlsx]
Columns: Product Name | EAN | SKU | Images (or "Image Link") | Colour (optional, overrides the guess)
"""
import json, re, sys
from pathlib import Path
from openpyxl import load_workbook

HERE = Path(__file__).parent

# name -> swatch hex. Multi-word names first so "dark blue" wins over "blue".
# ponytail: colour is guessed from the name/SKU text; add a "Colour" column when the guess falls short.
COLOURS = {
    "dark blue": "#1f2f5c", "light blue": "#9cc4e4", "baby pink": "#f4c2d0", "pista green": "#93c572",
    "dark green": "#1f5e3a", "rose red": "#c21e56", "multi color": "", "multi-color": "", "multicolor": "",
    "black": "#1b1c1f", "grey": "#8a8d91", "blue": "#2f5fa8", "white": "#f4f4f2", "cream": "#ece2cc",
    "pink": "#e8a0bf", "purple": "#6b3fa0", "green": "#3a8f4a", "cyan": "#3ec1d3", "teal": "#1f8a8a",
    "red": "#c62828", "brown": "#6e4b2a", "peach": "#f7c6a3", "lavender": "#b7a3d9", "silver": "#c0c4c8",
    "gold": "#d4af37", "golden": "#d4af37",  # not "transparent"/"clear": they describe windows and displays, not the item
}
IN_NAME = re.compile(r"\b(" + "|".join(COLOURS) + r")\b", re.I)


def colour_of(name, sku):
    found = IN_NAME.findall(name)
    if found:
        return found[-1].lower()  # variant colour is the last one mentioned
    return next((c for c in COLOURS if sku.lower().endswith(c.replace(" ", ""))), "")


def ean_key(ean):
    # Excel stores EANs as numbers, which drops leading zeros - so drop them everywhere.
    if isinstance(ean, float):
        ean = int(ean)
    return re.sub(r"\D", "", str(ean)).lstrip("0")


def image_url(link):
    drive = re.search(r"/d/([\w-]+)", link or "")
    return f"https://drive.google.com/thumbnail?id={drive[1]}&sz=w1000" if drive else (link or "")


def build(xlsx):
    rows = load_workbook(xlsx, read_only=True).active.iter_rows(values_only=True)
    col = {str(h).strip().lower(): i for i, h in enumerate(next(rows)) if h}
    products = {}
    for row in rows:
        get = lambda h: str(row[col[h]] or "").strip() if h in col else ""
        key = ean_key(row[col["ean"]] or "")
        if not key:
            continue
        if key in products:
            print(f"warning: duplicate EAN {key}, keeping the last row")
        colour = (get("colour") or get("color")).lower() or colour_of(get("product name"), get("sku"))
        products[key] = {
            "name": get("product name"), "sku": get("sku"), "colour": colour,
            "hex": COLOURS.get(colour, ""), "image": image_url(get("images") or get("image link") or get("image")),
        }
    out = HERE / "products.js"
    out.write_text("const PRODUCTS = " + json.dumps(products, indent=1, ensure_ascii=False) + ";\n", encoding="utf-8")
    print(f"{len(products)} products -> {out.name}")


if __name__ == "__main__":
    assert colour_of("Organizer Bag (Dark Blue)", "x") == "dark blue"
    assert colour_of("Black zipper bag - Grey", "x") == "grey"
    assert colour_of("Jewellery Organiser", "sqboxcream") == "cream"
    assert colour_of("Foldable Bag (Rose Red)", "fbag") == "rose red"
    assert colour_of("Jewellery Organiser", "jcasegolden") == "golden"
    assert ean_key(8906209350897.0) == ean_key("0 8906209350897") == "8906209350897"
    build(sys.argv[1] if len(sys.argv) > 1 else HERE / "products.xlsx")

"""Label PDFs -> labels/*.png + labels.js, so the app reads ASIN labels and shows each product's labels.
Run when the labels change:  python labels.py <labels folder>      (needs Node, to run zbar.js)
The folder holds one subfolder per label set, and the sheet that lists each set's SKUs ("lable and sku, images.xlsx").
"""
import json, re, subprocess, sys, tempfile
from pathlib import Path
import pymupdf
from openpyxl import load_workbook

HERE = Path(__file__).parent
SETS = {"3x2 inch lables": "E-trade", "3x2 lable Bigbasket": "Bigbasket", "3x2 lable UAE LABELS PRODUCT WISE": "UAE"}  # folder -> caption

# The printed text is no guide to the bars: some labels print the ASIN while their barcode holds the EAN.
# So read the bars, with the app's own decoder (zbar.js), and the import sees exactly what the phones will.
DECODE = """
const fs = require('fs'), [zbar, list] = process.argv.slice(1);
global.__dirname = require('path').dirname(zbar);
require('vm').runInThisContext(fs.readFileSync(zbar, 'utf8'));
(async () => {
  const out = [];
  for (const [file, width, height] of JSON.parse(fs.readFileSync(list, 'utf8')))
    out.push((await zbarWasm.scanGrayBuffer(Uint8Array.from(fs.readFileSync(file)).buffer, width, height)).map(s => s.decode()));
  console.log(JSON.stringify(out));
})();
"""


def barcodes(pixmaps):
    """The barcode values on each grey pixmap."""
    with tempfile.TemporaryDirectory() as tmp:
        pages = []
        for i, pix in enumerate(pixmaps):
            (Path(tmp) / f"{i}.raw").write_bytes(pix.samples)
            pages.append([str(Path(tmp) / f"{i}.raw"), pix.width, pix.height])
        (Path(tmp) / "pages.json").write_text(json.dumps(pages))
        run = subprocess.run(["node", "-e", DECODE, str(HERE / "zbar.js"), str(Path(tmp) / "pages.json")],
                             capture_output=True, text=True, check=True)
        return json.loads(run.stdout)


def norm(sku):
    return re.sub(r"[\W_]", "", sku.lower())  # labels say "fstool-pack-2", products.xlsx says "fstool-pack2"


def sheet_skus(xlsx):
    found = set()
    for ws in load_workbook(xlsx, read_only=True, data_only=True).worksheets:
        col = None
        for row in ws.iter_rows(values_only=True):
            cells = [str(c or "").strip().lower() for c in row]
            if col is not None and col < len(cells) and cells[col]:
                found.add(cells[col])
            elif col is None and "sku" in cells:
                col = cells.index("sku")
    return found


def sku_on(label, known):
    """The one known SKU written on the label or in its file name, else None."""
    hits = {s for s in known if re.search(rf"(?<![\w-]){re.escape(s)}(?![\w-])", label)}
    return hits.pop() if len(hits) == 1 else None


def main(root):
    products = json.loads((HERE / "products.js").read_text(encoding="utf-8").split("=", 1)[1].rstrip().rstrip(";"))
    product_sku = {norm(p["sku"]): p["sku"].lower() for p in products.values()}
    known = set(product_sku.values()).union(*(sheet_skus(x) for x in root.glob("*.xlsx")))
    (HERE / "labels").mkdir(exist_ok=True)
    for old in (HERE / "labels").glob("*.png"):
        old.unlink()
    pages = []
    for pdf in sorted(root.rglob("*.pdf")):
        group = SETS.get(pdf.relative_to(root).parts[0], pdf.relative_to(root).parts[0])
        for page in pymupdf.open(pdf):
            sku = sku_on(f"{pdf.stem} {page.get_text()}".lower(), known)
            pages.append((pdf, page.number, group, sku, page.get_pixmap(dpi=200, colorspace=pymupdf.csGRAY)))
    labels = {}
    for (pdf, number, group, sku, pix), codes in zip(pages, barcodes([p[-1] for p in pages])):
        if len(codes) != 1 or not sku:
            print(f"skipped {pdf.name} page {number + 1}: barcodes {codes or 'none'}, SKU {sku or 'not found'}")
            continue
        src = f"labels/{norm(group)}-{codes[0]}.png"
        if src not in labels:  # most labels are both in their own PDF and in the combined one
            pix.save(HERE / src)
            labels[src] = {"code": codes[0], "sku": product_sku.get(norm(sku), sku), "set": group, "src": src}
    for label in labels.values():
        if norm(label["sku"]) not in product_sku:
            print(f"no product for SKU {label['sku']} ({label['set']} label {label['code']}): add it to products.xlsx")
    (HERE / "labels.js").write_text("const LABELS = " + json.dumps(list(labels.values()), indent=1) + ";\n", encoding="utf-8")
    print(f"{len(labels)} labels -> labels/, labels.js")


if __name__ == "__main__":
    assert sku_on("lazer stool, pack of 2 | fstool-pack-2 b0ghr7kjq6", {"fstool", "fstool-pack-2"}) == "fstool-pack-2"
    assert sku_on("pcblack label uae", {"pcblack", "pc"}) == "pcblack"
    assert norm("fstool-pack-2") == norm("fstool-Pack2")
    main(Path(sys.argv[1]))
